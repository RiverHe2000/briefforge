"""Export one immutable research snapshot to editable Word and PowerPoint.

No model, network, database, active sources or evaluation references are consulted.
PptxGenJS is invoked without a shell and receives only the snapshot projection.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from .export_contract import EXPORT_RENDERER_VERSION

STATUS_LABELS = {"supported": "有来源支持", "uncertain": "待确认", "contradicted": "来源存在冲突"}


def _runtime_root() -> Path:
    return Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"


def build_export_model(report: dict[str, Any]) -> dict[str, Any]:
    """Validate provenance and project identical facts for both formats."""
    snapshot = json.loads(json.dumps(report, ensure_ascii=False))
    sources = {source["id"]: source for source in snapshot.get("sources", [])}
    claims = snapshot.get("claims", [])
    known_ids = {claim["id"] for claim in claims}
    if len(known_ids) != len(claims):
        raise ValueError("报告包含重复事实 ID，无法导出。")
    if not snapshot.get("synthetic") and any(source.get("synthetic") for source in sources.values()):
        raise ValueError("公开资料报告不能隐藏合成来源标签。")
    for claim in claims:
        if claim.get("status") not in STATUS_LABELS:
            raise ValueError("报告包含未知事实状态。")
        for evidence in claim.get("evidence", []):
            source = sources.get(evidence.get("source_id"))
            quote = evidence.get("quote", "")
            if source is None or not quote or quote not in source.get("text", ""):
                raise ValueError("事实引用不属于此报告的冻结资料版本。")
        if claim.get("status") == "supported" and not claim.get("evidence"):
            raise ValueError("受支持事实缺少来源，无法导出。")
    for section in [*snapshot.get("sections", []), *snapshot.get("comparison", [])]:
        if not set(section.get("claim_ids", [])).issubset(known_ids):
            raise ValueError("报告内容引用了未知事实 ID。")
    source_numbers = {source_id: i + 1 for i, source_id in enumerate(sources)}
    for claim in claims:
        claim["status_label"] = STATUS_LABELS[claim["status"]]
        claim["citations"] = sorted({source_numbers[e["source_id"]] for e in claim.get("evidence", [])})
    snapshot["source_numbers"] = source_numbers
    snapshot["disclosure"] = "虚构测试资料" if snapshot.get("synthetic") else "公开资料研究"
    snapshot["mode_label"] = "固定响应回放" if snapshot.get("mode") == "replay" else "真实模型研究草稿"
    snapshot["fingerprint"] = (
        snapshot.get("content_hash")
        or hashlib.sha256(json.dumps(report, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    )
    snapshot["chart"] = _comparison_chart(snapshot)
    snapshot["manifest"] = {
        "exporter_version": EXPORT_RENDERER_VERSION,
        "report_id": snapshot.get("id"),
        "report_version": snapshot.get("version"),
        "content_hash": snapshot["fingerprint"],
        "synthetic": bool(snapshot.get("synthetic")),
        "mode": snapshot.get("mode"),
        "model_profile": snapshot.get("model_profile", "qwen-default"),
        "sections": snapshot.get("sections", []),
        "claims": [
            {
                key: claim.get(key)
                for key in (
                    "id",
                    "subject",
                    "dimension",
                    "statement",
                    "status",
                    "value",
                    "conditions",
                    "evidence",
                    "computation",
                )
            }
            for claim in claims
        ],
        "sources": [
            {
                key: source.get(key)
                for key in ("id", "title", "version", "sha256", "published_at", "url", "synthetic")
            }
            for source in sources.values()
        ],
    }
    return snapshot


def _comparison_chart(report: dict[str, Any]) -> dict[str, Any]:
    points = []
    for row in report.get("comparison", []):
        text = str(row.get("price", ""))
        match = re.search(r"\b(USD|AUD|EUR|GBP)\s*([0-9]+(?:\.[0-9]+)?)\s*/\s*席位\s*/\s*月", text)
        if not match or "按年付费" not in text or "未含税" not in text:
            points = []
            break
        # Do not turn an uncertain or unsupported price into a chart fact.
        price_claims = [
            c
            for c in report.get("claims", [])
            if c.get("subject") == row.get("competitor") and c.get("dimension") == "价格"
        ]
        if not price_claims or any(c.get("status") != "supported" for c in price_claims):
            points = []
            break
        points.append((row["competitor"], float(match.group(2)), match.group(1)))
    if len(points) >= 2 and len({p[2] for p in points}) == 1:
        currency = points[0][2]
        return {
            "kind": "price",
            "title": "团队版年付月均报价",
            "series": f"{currency}/席位/月",
            "labels": [p[0] for p in points],
            "values": [p[1] for p in points],
            "note": f"{currency}/席位/月，按年付费，未含税。仅比较资料中的团队版标价，企业版和月付条件见对应说明。",
        }
    labels = ["有来源支持", "待确认", "来源冲突"]
    return {
        "kind": "coverage",
        "title": "报告事实状态",
        "series": "事实条数",
        "labels": labels,
        "values": [
            sum(c.get("status") == status for c in report.get("claims", []))
            for status in ("supported", "uncertain", "contradicted")
        ],
        "note": "按此报告保存的事实状态计数。此图不衡量产品质量或真实市场表现。",
    }


def export_report(report: dict[str, Any], format: str, output_dir: Path) -> Path:
    if format not in {"docx", "pptx"}:
        raise ValueError("支持的导出格式为 docx 和 pptx。")
    model = build_export_model(report)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_id = re.sub(r"[^A-Za-z0-9_-]", "", str(report.get("id", "")))[:80] or model["fingerprint"][:16]
    destination = output_dir / f"briefforge-{safe_id}-v{int(report.get('version', 1))}.{format}"
    with tempfile.TemporaryDirectory(prefix=".export-", dir=output_dir) as working:
        draft = Path(working) / f"report.{format}"
        if format == "docx":
            _export_docx(model, draft)
        else:
            _export_pptx(model, draft)
        if not draft.is_file() or draft.stat().st_size < 500:
            raise RuntimeError("办公文件生成失败，未产生可用文件。")
        os.replace(draft, destination)
    return destination


def _set_font(style, size: float, bold: bool = False) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    style.font.name = "Aptos"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _table(document, headers: list[str], rows: list[list[str]], widths: list[float]):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for row_index, data in enumerate([headers, *rows]):
        row = table.rows[0] if row_index == 0 else table.add_row()
        for cell, text, width in zip(row.cells, data, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(text)
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{side}")
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    border.set(qn(f"w:{key}"), value)
                borders.append(border)
            tcpr.append(borders)
            shade = OxmlElement("w:shd")
            shade.set(
                qn("w:fill"), "253D4B" if row_index == 0 else ("F0F3F4" if row_index % 2 == 0 else "FFFFFF")
            )
            tcpr.append(shade)
            margins = OxmlElement("w:tcMar")
            for name in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{name}")
                margin.set(qn("w:w"), "85")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            tcpr.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _export_docx(report: dict[str, Any], destination: Path) -> None:
    from xml.sax.saxutils import escape

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    from .export_support.word_chart import add_editable_chart

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.72), Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.78)
    section.header_distance, section.footer_distance = Inches(0.3), Inches(0.3)
    for name, size, bold in [
        ("Normal", 11, False),
        ("Title", 27, True),
        ("Subtitle", 12, False),
        ("Heading 1", 19, True),
        ("Heading 2", 13, True),
        ("Caption", 9, False),
    ]:
        _set_font(document.styles[name], size, bold)
        ppr = document.styles[name].element.find(qn("w:pPr"))
        if ppr is not None:
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
    normal = document.styles["Normal"].paragraph_format
    normal.space_after, normal.line_spacing = Pt(7), Pt(15)
    # Disable East Asian document-grid snapping; it otherwise expands caption rows.
    for style_name in ("Normal", "Caption", "Heading 1", "Heading 2"):
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        document.styles[style_name].element.get_or_add_pPr().append(snap)
    document.styles["Caption"].paragraph_format.line_spacing = Pt(12)
    document.styles["Caption"].paragraph_format.space_after = Pt(5)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(5)
    document.core_properties.title = report.get("title", "行业与竞品研究")
    document.core_properties.author = "BriefForge"
    document.core_properties.subject = f"{report['disclosure']}；报告版本 {report.get('version', 1)}"
    document.core_properties.keywords = f"report:{report.get('id')};sha256:{report['fingerprint']}"
    header = section.header.paragraphs[0]
    header.text = f"BriefForge    {report['disclosure']}"
    header.style = document.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f"{report['disclosure']}    报告 v{report.get('version', 1)}    ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    document.add_paragraph("行业与竞品研究报告", "Title")
    document.add_paragraph(report.get("title", ""), "Subtitle")
    document.add_paragraph(
        f"{report['disclosure']}    {report['mode_label']}    版本 {report.get('version', 1)}"
    )
    document.add_paragraph(f"报告日期 {str(report.get('created_at', ''))[:10]}")
    document.add_heading("研究结论", 1)
    document.add_paragraph(
        report.get("executive_summary") or "此报告尚未形成摘要，以下仅列出已有事实与待确认事项。"
    )
    document.add_heading("阅读范围", 2)
    document.add_paragraph(
        "比较结论保留资料日期、套餐、计费周期与证据限制。价格不等于实际合同报价，未披露的能力保留为待确认。"
    )
    if report.get("synthetic"):
        document.add_paragraph(
            "本文使用虚构厂商与合成资料演示研究流程，不能用于真实采购、商业决策或市场规模判断。"
        )
    document.add_paragraph(
        f"冻结资料 {len(report.get('sources', []))} 份，事实 {len(report.get('claims', []))} 条。完整来源目录见文末。",
        "Caption",
    )
    document.add_page_break()
    document.add_heading("竞品比较", 1)
    rows = [
        [r.get("competitor", ""), r.get("price", "未形成价格结论"), r.get("sso", "未形成 SSO 结论")]
        for r in report.get("comparison", [])
    ]
    if rows:
        _table(document, ["产品", "价格及条件", "SSO 及适用范围"], rows, [1.22, 2.85, 2.85])
    else:
        document.add_paragraph("当前报告没有可供比较的竞品记录。")
    document.add_heading(report["chart"]["title"], 2)
    add_editable_chart(document, report["chart"])
    document.add_paragraph(report["chart"]["note"], "Caption")
    subjects = list(dict.fromkeys(r.get("competitor", "") for r in report.get("comparison", [])))
    if not subjects:
        subjects = list(dict.fromkeys(c.get("subject", "") for c in report.get("claims", [])))
    for subject in subjects:
        document.add_page_break()
        document.add_heading(subject or "其他研究事实", 1)
        # A profile presents the canonical facts directly; the prose synthesis is
        # retained in the embedded snapshot rather than repeating each fact twice.
        for claim in report.get("claims", []):
            if claim.get("subject", "") != subject:
                continue
            document.add_heading(claim.get("dimension") or "研究事实", 2)
            document.add_paragraph(claim.get("statement", ""))
            if claim.get("conditions"):
                document.add_paragraph("适用条件：" + "；".join(claim["conditions"]), "Caption")
            refs = " ".join(f"[{n:02d}]" for n in claim["citations"])
            document.add_paragraph(f"{claim['status_label']}。来源 {refs or '无直接引用'}", "Caption")
    if report.get("sections"):
        document.add_page_break()
        document.add_heading("研究发现与待确认事项", 1)
        for item in report["sections"]:
            if any(subject and subject in item.get("heading", "") for subject in subjects):
                continue
            document.add_heading(item.get("heading") or "研究发现", 2)
            document.add_paragraph(item.get("body", ""))
        for claim in report.get("claims", []):
            if claim.get("subject", "") not in subjects:
                document.add_paragraph(
                    f"{claim.get('subject', '')}  {claim.get('dimension', '')}：{claim.get('statement', '')}"
                )
                if claim.get("conditions"):
                    document.add_paragraph("适用条件：" + "；".join(claim["conditions"]), "Caption")
        for item in report.get("unresolved", []):
            document.add_paragraph(str(item), "List Bullet")
        if report.get("changes"):
            document.add_heading("本版本变化", 2)
            for item in report["changes"]:
                document.add_paragraph(str(item))
    else:
        document.add_page_break()
        document.add_heading("待确认事项", 1)
        for item in report.get("unresolved") or [
            "本报告没有列出额外待确认事项；事实状态与引用仍需结合对应资料理解。"
        ]:
            document.add_paragraph(str(item))
    sources = report.get("sources", [])
    for start in range(0, max(1, len(sources)), 15):
        document.add_page_break()
        document.add_heading("资料来源" if start == 0 else "资料来源续", 1)
        for source in sources[start : start + 15]:
            number = report["source_numbers"][source["id"]]
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = Pt(12)
            paragraph.add_run(f"[{number:02d}] {source['title']}").bold = True
            paragraph.add_run(
                f"\n资料日期 {str(source.get('published_at') or '未注明')[:10]}    版本 {source.get('version', 1)}"
            )
            if source.get("url"):
                paragraph.add_run(f"\n{source['url']}")
            for run in paragraph.runs:
                run.font.size = Pt(9)
        if not sources:
            document.add_paragraph("此报告没有冻结来源。")
    # Embed all facts, quotes and conditions for machine checking and later audit.
    payload = json.dumps(report["manifest"], ensure_ascii=False, sort_keys=True)
    part = Part(
        PackURI("/customXml/briefforge-snapshot.xml"),
        "application/xml",
        f'<snapshot xmlns="urn:briefforge:snapshot">{escape(payload)}</snapshot>'.encode(),
        document.part.package,
    )
    document.part.relate_to(
        part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
    )
    document.save(destination)


def _export_pptx(report: dict[str, Any], destination: Path) -> None:
    root = _runtime_root()
    configured_node = os.environ.get("BRIEFFORGE_NODE")
    node = configured_node or (
        str(root / "node/bin/node.exe") if (root / "node/bin/node.exe").exists() else shutil.which("node")
    )
    if not node:
        raise RuntimeError("PPT 导出需要 Node.js 和 PptxGenJS，请配置 BRIEFFORGE_NODE。")
    environment = os.environ.copy()
    modules = os.environ.get("BRIEFFORGE_NODE_MODULES") or str(root / "node/node_modules")
    if Path(modules).is_dir():
        environment["NODE_PATH"] = modules + (
            os.pathsep + environment["NODE_PATH"] if environment.get("NODE_PATH") else ""
        )
    input_path = destination.with_suffix(".json")
    input_path.write_text(json.dumps(report, ensure_ascii=False), encoding="utf-8")
    script = Path(__file__).parent / "export_support" / "build_deck.cjs"
    result = subprocess.run(
        [str(node), str(script), str(input_path), str(destination)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=90,
        env=environment,
        check=False,
    )
    if result.returncode:
        message = (result.stderr or result.stdout)[-1500:]
        raise RuntimeError(f"PPT 导出未完成：{message}")
    _normalize_pptx_content_types(destination)


def _normalize_pptx_content_types(path: Path) -> None:
    """Normalize generator package identifiers while preserving native content.

    Some packaged versions list one master per slide but write one shared master.
    PowerPoint rejects such packages. The generator can also reuse nonvisual IDs
    when a native table or chart follows text. This template has no animations or
    connectors referencing those IDs; sequential reassignment is unambiguous.
    """
    from xml.etree import ElementTree as ET
    from zipfile import ZipFile

    with ZipFile(path) as archive:
        parts = {item.filename: (item, archive.read(item.filename)) for item in archive.infolist()}
    namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    root = ET.fromstring(parts["[Content_Types].xml"][1])
    changed = False
    for element in list(root):
        if (
            element.tag == f"{{{namespace}}}Override"
            and element.attrib.get("PartName", "").lstrip("/") not in parts
        ):
            root.remove(element)
            changed = True
    presentation_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    for name, (info, content) in list(parts.items()):
        if not re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
            continue
        slide = ET.fromstring(content)
        elements = slide.findall(f".//{{{presentation_ns}}}cNvPr")
        ids = [element.get("id") for element in elements]
        if len(ids) != len(set(ids)):
            # String substitution avoids changing namespace prefixes throughout
            # a slide, particularly mc:Ignorable extension declarations.
            counter = iter(range(1, len(elements) + 1))
            updated = re.sub(
                rb'(<p:cNvPr\s+id=")[^"]+("[^>]*>)',
                lambda match, counter=counter: match.group(1) + str(next(counter)).encode() + match.group(2),
                content,
            )
            parts[name] = (info, updated)
            changed = True
    if not changed:
        return
    ET.register_namespace("", namespace)
    info, _ = parts["[Content_Types].xml"]
    parts["[Content_Types].xml"] = (info, ET.tostring(root, encoding="utf-8", xml_declaration=True))
    normalized = path.with_suffix(".normalized.pptx")
    with ZipFile(normalized, "w") as archive:
        for info, data in parts.values():
            archive.writestr(info, data)
    os.replace(normalized, path)
