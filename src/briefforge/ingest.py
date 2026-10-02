"""Bounded file extraction and public-web fetches. Input documents are data, never instructions."""

from __future__ import annotations

import asyncio
import csv
import hashlib
import http.client
import io
import ipaddress
import re
import socket
import ssl
import zipfile
from pathlib import Path
from urllib.parse import urljoin, urlsplit

from bs4 import BeautifulSoup

MAX_FILE = 10 * 1024 * 1024
MAX_TEXT = 500_000


def _zip_check(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 3000 or sum(e.file_size for e in entries) > 40 * 1024 * 1024:
            raise ValueError("文件解压后过大，最多40MB")
        if any(e.flag_bits & 1 for e in entries):
            raise ValueError("不支持加密文档")


def decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError("无法识别文本编码，请保存为UTF-8")


def html_text(raw: str) -> tuple[str, str, str | None]:
    soup = BeautifulSoup(raw, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else "网页资料"
    date = None
    for tag in soup.select('meta[property="article:published_time"], meta[name="date"], time[datetime]'):
        candidate = tag.get("content") or tag.get("datetime")
        if candidate and re.match(r"^\d{4}-\d{2}-\d{2}", str(candidate)):
            date = str(candidate)[:35]
            break
    for node in soup(["script", "style", "noscript", "svg", "nav", "form", "iframe"]):
        node.decompose()
    text = "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
    return title[:240], text[:MAX_TEXT], date


def parse_file(filename: str, raw: bytes) -> dict:
    if not raw:
        raise ValueError("文件为空")
    if len(raw) > MAX_FILE:
        raise ValueError("单个文件不能超过10MB")
    ext = Path(filename).suffix.lower()
    title = Path(filename).name[:240]
    if ext in {".txt", ".md"}:
        text = decode(raw)
    elif ext == ".csv":
        rows = list(csv.reader(io.StringIO(decode(raw))))
        if len(rows) > 10000:
            raise ValueError("CSV最多10000行")
        text = "\n".join(f"第{i + 1}行：" + " | ".join(row) for i, row in enumerate(rows))
    elif ext in {".html", ".htm"}:
        title, text, _ = html_text(decode(raw))
    elif ext == ".docx":
        from docx import Document

        _zip_check(raw)
        doc = Document(io.BytesIO(raw))
        lines = [f"段落{i + 1}：{p.text}" for i, p in enumerate(doc.paragraphs) if p.text.strip()]
        for index, table in enumerate(doc.tables):
            lines.extend(
                f"表{index + 1}第{rowidx + 1}行：" + " | ".join(c.text for c in row.cells)
                for rowidx, row in enumerate(table.rows)
            )
        text = "\n".join(lines)
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        _zip_check(raw)
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        lines = []
        try:
            for sheet in workbook:
                if sheet.max_row > 10000 or sheet.max_column > 100:
                    raise ValueError("工作表最多10000行、100列")
                for idx, row in enumerate(sheet.iter_rows(values_only=True)):
                    cells = [str(value) if value is not None else "" for value in row]
                    if any(cells):
                        lines.append(f"{sheet.title} 第{idx + 1}行：" + " | ".join(cells))
        finally:
            workbook.close()
        text = "\n".join(lines)
    elif ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("不支持加密PDF")
        if len(reader.pages) > 100:
            raise ValueError("PDF最多100页")
        text = "\n".join(f"第{i + 1}页\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages))
        if len(re.sub(r"第\d+页|\s", "", text)) < 5:
            raise ValueError("PDF没有可提取文本；当前版本不支持扫描件OCR")
    else:
        raise ValueError("支持 CSV、XLSX、DOCX、文本型PDF、HTML、TXT、Markdown")
    if not text.strip():
        raise ValueError("未提取到可读取内容")
    if len(text) > MAX_TEXT:
        raise ValueError("提取文本超过50万字符，请拆分资料")
    return {"title": title, "text": text, "kind": ext.lstrip("."), "sha256": hashlib.sha256(raw).hexdigest()}


def public_addresses(url: str) -> tuple[str, int, list[str]]:
    u = urlsplit(url)
    if u.scheme not in {"http", "https"} or not u.hostname or u.username or u.password:
        raise ValueError("仅支持公开HTTP或HTTPS网址")
    port = u.port or (443 if u.scheme == "https" else 80)
    if port not in {80, 443}:
        raise ValueError("仅允许标准网页端口")
    host = u.hostname.encode("idna").decode("ascii")
    addresses = list(
        dict.fromkeys(info[4][0] for info in socket.getaddrinfo(host, port, type=socket.SOCK_STREAM))
    )
    if not addresses or any(not ipaddress.ip_address(ip).is_global for ip in addresses):
        raise ValueError("不能读取本机、局域网或保留地址")
    return host, port, addresses


def _fetch_public(url: str) -> dict:
    for _ in range(4):
        host, port, addresses = public_addresses(url)
        u = urlsplit(url)
        connection = http.client.HTTPConnection(host, port, timeout=12)
        # Connect to the validated IP itself. DNS cannot change between validation
        # and the actual connection; TLS still verifies the original hostname.
        sock = socket.create_connection((addresses[0], port), timeout=12)
        if u.scheme == "https":
            sock = ssl.create_default_context().wrap_socket(sock, server_hostname=host)
        connection.sock = sock
        try:
            path = u.path or "/"
            if u.query:
                path += "?" + u.query
            connection.request(
                "GET",
                path,
                headers={
                    "Host": host,
                    "User-Agent": "BriefForge/0.1 Research (+local user request)",
                    "Accept": "text/html,text/plain,application/pdf",
                    "Accept-Encoding": "identity",
                },
            )
            response = connection.getresponse()
            if response.status in {301, 302, 303, 307, 308}:
                location = response.getheader("Location")
                if not location:
                    raise ValueError("网页重定向缺少地址")
                url = urljoin(url, location)
                continue
            if response.status != 200:
                raise ValueError(f"网页返回HTTP {response.status}")
            raw = response.read(MAX_FILE + 1)
            if len(raw) > MAX_FILE:
                raise ValueError("网页超过10MB")
            content_type = response.getheader("Content-Type", "").lower()
            if "pdf" in content_type:
                data = parse_file("source.pdf", raw)
                title, text, published = data["title"], data["text"], None
            elif "text/html" in content_type or "application/xhtml" in content_type:
                charset = re.search(r"charset=([\w-]+)", content_type)
                try:
                    body = raw.decode(charset.group(1) if charset else "utf-8", errors="replace")
                except LookupError:
                    body = raw.decode("utf-8", errors="replace")
                title, text, published = html_text(body)
            elif "text/plain" in content_type:
                title, text, published = host, decode(raw), None
            else:
                raise ValueError("网址没有返回可读取的网页、文本或PDF")
            validate_page_body(title, text)
            return {
                "title": title,
                "text": text[:MAX_TEXT],
                "url": url,
                "published_at": published,
                "kind": "web",
            }
        finally:
            connection.close()
    raise ValueError("网页重定向次数过多")


def validate_page_body(title: str, text: str) -> None:
    if len(text.strip()) < 30:
        raise ValueError("网页正文不足，可能需要登录或动态加载")
    if len(text) < 1000 and re.search(
        r"CSS Error\s+Refresh|Sorry to interrupt\s+CSS Error", text, re.IGNORECASE
    ):
        raise ValueError("网页只返回加载错误提示，未取得可用正文")
    # A successful HTTP response can still be a login or browser challenge.
    # Reject short interstitials, retaining the URL and reason in the run event.
    if (
        len(text) < 2000
        and re.search(
            r"login to access|sign in to continue|verify you are human|your browser was unable to load",
            text,
            re.IGNORECASE,
        )
        and re.search(r"login|sign.in|access.denied|just a moment|验证|登录", title, re.IGNORECASE)
    ):
        raise ValueError("网页返回登录或访问验证页面，未取得可用正文")


async def fetch_public_url(url: str) -> dict:
    return await asyncio.to_thread(_fetch_public, url)
