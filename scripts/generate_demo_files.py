"""Generate 30 uploadable, explicitly fictional fixtures from demo_sources only.

No evaluation references or answer files are read. The manifest preserves exact
source text and metadata, while every material is verified with the application's
real parse_file(filename, raw) ingestion function before success is reported.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import math
import re
import sys
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from briefforge.demo import DISCLAIMER, demo_sources
from briefforge.ingest import parse_file

FORMATS = ("txt", "md", "html", "csv", "xlsx", "docx")


def sha256(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def metadata_rows(source: dict) -> list[tuple[str, str]]:
    return [
        ("资料标识", DISCLAIMER),
        ("原始标题", source["title"]),
        ("资料归属", source.get("competitor") or "行业背景"),
        ("发布日期", source.get("published_at") or "未披露"),
        ("资料类型", source["kind"]),
        ("资料编号", source["logical_key"]),
        ("来源链接", source.get("url") or "无 这是虚构测试材料"),
    ]


def source_lines(source: dict) -> list[str]:
    return [line for line in source["text"].splitlines() if line.strip()]


def write_material(path: Path, source: dict) -> None:
    rows = metadata_rows(source)
    if path.suffix == ".txt":
        header = "\n".join(f"{name}：{value}" for name, value in rows)
        path.write_text(header + "\n\n资料原文\n\n" + source["text"] + "\n", encoding="utf-8")
    elif path.suffix == ".md":
        header = "\n".join(f"- {name}：{value}" for name, value in rows)
        path.write_text(f"# {source['title']}\n\n{header}\n\n## 资料原文\n\n{source['text']}\n", encoding="utf-8")
    elif path.suffix == ".html":
        description = "".join(f"<dt>{html.escape(name)}</dt><dd>{html.escape(value)}</dd>" for name, value in rows)
        paragraphs = "\n".join(f"<p>{html.escape(line)}</p>" for line in source_lines(source))
        title = html.escape(source["title"])
        page = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta property="article:published_time" content="{html.escape(source.get('published_at') or '')}">
<title>{title}</title><style>body{{font:16px/1.9 'Microsoft YaHei',sans-serif;color:#253b2f;background:#f5f7f1;margin:0;padding:36px}}article{{max-width:850px;margin:auto;background:white;padding:36px 44px;border:1px solid #dfe6d7}}h1{{font-size:25px;line-height:1.6}}h2{{font-size:19px;margin-top:30px}}dl{{display:grid;grid-template-columns:110px 1fr;gap:10px;font-size:13px}}dt{{color:#607356}}dd{{margin:0}}p{{white-space:pre-wrap}}</style></head>
<body><article><h1>{title}</h1><dl>{description}</dl><h2>资料原文</h2>{paragraphs}</article></body></html>"""
        path.write_text(page, encoding="utf-8")
    elif path.suffix == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["字段", "资料内容"])
            writer.writerows(rows)
            writer.writerows((f"原文段落 {index}", text) for index, text in enumerate(source_lines(source), 1))
    elif path.suffix == ".xlsx":
        write_workbook(path, source, rows)
    elif path.suffix == ".docx":
        write_document(path, source, rows)
    else:
        raise ValueError(f"Unsupported fixture format: {path.suffix}")


def write_workbook(path: Path, source: dict, rows: list[tuple[str, str]]) -> None:
    # The reusable Python generator deliberately uses the project's existing
    # openpyxl dependency, keeping fixture generation free of extra toolchains.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    book = Workbook()
    sheet = book.active
    sheet.title = "虚构研究资料"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 100
    sheet.append(["字段", "资料内容"])
    for row in rows:
        sheet.append(row)
    for index, text in enumerate(source_lines(source), 1):
        sheet.append([f"原文段落 {index}", text])
    for row in sheet:
        for cell in row:
            cell.data_type = "s"  # Source prose is literal text, never a spreadsheet formula.
            cell.font = Font(name="Microsoft YaHei", size=11, color="24372C")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="E5EADB"))
        content = str(row[1].value or "")
        display_units = sum(2 if ord(character) > 255 else 1 for character in content)
        sheet.row_dimensions[row[0].row].height = max(30, math.ceil(display_units / 96) * 18 + 15)
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244936")
        cell.alignment = Alignment(vertical="center")
    for row in range(2, len(rows) + 2):
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor="F0F4E8")
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f"A1:B{sheet.max_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    book.properties.title = source["title"]
    book.properties.subject = "虚构测试资料"
    book.properties.creator = "BriefForge"
    book.save(path)
    book.close()


def write_document(path: Path, source: dict, rows: list[tuple[str, str]]) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(22)
    section.left_margin = section.right_margin = Mm(24)
    for name in ("Normal", "Title", "Heading 1"):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(8)
    document.styles["Title"].font.size = Pt(19)
    document.styles["Heading 1"].font.size = Pt(13)
    document.add_paragraph(source["title"], "Title")
    document.add_paragraph(DISCLAIMER)
    for name, value in rows[2:]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(f"{name}：").bold = True
        paragraph.add_run(value)
    document.add_paragraph("资料原文", "Heading 1")
    for text in source_lines(source):
        paragraph = document.add_paragraph(text)
        # Prevent a single paragraph splitting across a page while retaining natural flow.
        keep = OxmlElement("w:keepLines")
        paragraph._p.get_or_add_pPr().append(keep)
    document.core_properties.title = source["title"]
    document.core_properties.subject = "虚构测试资料"
    document.core_properties.author = "BriefForge"
    document.core_properties.created = datetime(2026, 10, 2, tzinfo=UTC)
    document.core_properties.modified = datetime(2026, 10, 2, tzinfo=UTC)
    document.save(path)


def validate_material(path: Path, source: dict) -> dict:
    raw = path.read_bytes()
    extracted = parse_file(path.name, raw)
    text = extracted["text"]
    assert DISCLAIMER in text, f"Missing synthetic disclosure: {path.name}"
    assert source["title"] in text, f"Missing source title: {path.name}"
    position = 0
    for line in source_lines(source):
        found = text.find(line, position)
        assert found >= 0, f"Original source paragraph lost or reordered in {path.name}: {line[:40]}"
        position = found + len(line)
    return {
        "filename": path.name, "format": path.suffix.lstrip("."), "status": "passed",
        "synthetic_disclosure_present": True, "original_nonempty_paragraphs": len(source_lines(source)),
        "all_original_paragraphs_preserved_in_order": True,
        "source_text_sha256": sha256(source["text"].encode("utf-8")),
        "file_sha256": sha256(raw), "size_bytes": len(raw), "extracted_characters": len(text),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--scenario", default="default")
    parser.add_argument("--output", type=Path, default=ROOT / "data" / "demo-files")
    args = parser.parse_args()
    output = args.output.resolve()
    sources = demo_sources(args.scenario)
    assert len(sources) == 30 and all(source.get("synthetic") for source in sources)
    output.mkdir(parents=True, exist_ok=True)
    filenames = [f"{index + 1:02d}-{re.sub(r'[^a-zA-Z0-9_.-]', '-', source['logical_key'])}.{FORMATS[index % len(FORMATS)]}" for index, source in enumerate(sources)]
    unexpected = [path.name for path in output.iterdir() if path.is_file() and path.suffix.lstrip(".") in FORMATS and path.name not in filenames]
    if unexpected:
        raise ValueError(f"Unexpected materials in output directory; inspect before regenerating: {unexpected}")
    materials, checks = [], []
    for filename, source in zip(filenames, sources, strict=True):
        path = output / filename
        write_material(path, source)
        check = validate_material(path, source)
        checks.append(check)
        materials.append({"filename": filename, "format": path.suffix.lstrip("."), "source": source, "source_text_sha256": check["source_text_sha256"], "file_sha256": check["file_sha256"]})
    format_counts = dict(Counter(check["format"] for check in checks))
    assert len(checks) == 30 and all(format_counts.get(format) == 5 for format in FORMATS)
    actual = [path for path in output.iterdir() if path.is_file() and path.suffix.lstrip(".") in FORMATS]
    assert len(actual) == 30
    manifest = {
        "synthetic": True, "disclosure": DISCLAIMER, "scenario": args.scenario,
        "material_count": 30, "format_counts": format_counts,
        "source_function": "briefforge.demo.demo_sources", "evaluation_answers_included": False,
        "instructions": "Upload materials through the source library. Set the matching competitor before upload. The manifest preserves canonical metadata; the generic file uploader extracts text and does not automatically load this manifest.",
        "materials": materials,
    }
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    validation = {
        "status": "passed", "material_count": 30, "format_counts": format_counts,
        "ingestion_function": "briefforge.ingest.parse_file(filename, raw)",
        "evaluation_answers_read": False, "model_calls": 0,
        "all_files_roundtripped": True, "all_files_preserve_disclosure_and_original_paragraphs": True,
        "verification_scope": "Content extraction and source paragraph preservation; Office page layout is not assessed by this check.",
        "checks": checks,
    }
    artifact = ROOT / "artifacts" / "demo-files-validation.json"
    artifact.parent.mkdir(parents=True, exist_ok=True)
    artifact.write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"PASS 30 fictional materials generated and roundtripped: {format_counts}")
    print(f"Materials: {output}")
    print(f"Validation: {artifact}")


if __name__ == "__main__":
    main()
