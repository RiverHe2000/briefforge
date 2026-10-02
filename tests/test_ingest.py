import io
import socket

import pytest
from docx import Document
from openpyxl import Workbook

from briefforge.ingest import MAX_FILE, html_text, parse_file, public_addresses, validate_page_body


def test_csv_preserves_rows_and_utf8():
    result = parse_file("价格.csv", "产品,单价,条件\nA,12,每席位每月".encode())
    assert "第2行：A | 12 | 每席位每月" in result["text"]


def test_docx_and_xlsx_locators():
    document = Document()
    document.add_paragraph("企业版支持SSO")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "USD"
    table.cell(0, 1).text = "20"
    stream = io.BytesIO()
    document.save(stream)
    text = parse_file("sample.docx", stream.getvalue())["text"]
    assert "段落1" in text and "表1第1行" in text
    workbook = Workbook()
    workbook.active.append(["套餐", "价格"])
    workbook.active.append(["企业版", 20])
    stream = io.BytesIO()
    workbook.save(stream)
    assert "第2行：企业版 | 20" in parse_file("p.xlsx", stream.getvalue())["text"]


def test_unsupported_and_oversize():
    with pytest.raises(ValueError):
        parse_file("macro.xlsm", b"content")
    with pytest.raises(ValueError):
        parse_file("x.txt", b"a" * (MAX_FILE + 1))


def test_html_removes_active_content_and_preserves_date():
    title, text, date = html_text(
        '<title>Pricing</title><meta property="article:published_time" content="2026-01-02"><script>secret</script><p>USD 20</p>'
    )
    assert title == "Pricing" and "secret" not in text and date == "2026-01-02"


@pytest.mark.parametrize("ip", ["127.0.0.1", "10.1.2.3", "169.254.169.254", "::1", "192.168.1.1", "0.0.0.0"])
def test_ssrf_denies_private_dns(monkeypatch, ip):
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(2, 1, 6, "", (ip, 443))])
    with pytest.raises(ValueError):
        public_addresses("https://example.test")


def test_ssrf_denies_mixed_dns_and_url_credentials(monkeypatch):
    monkeypatch.setattr(
        socket,
        "getaddrinfo",
        lambda *a, **k: [(2, 1, 6, "", ("93.184.216.34", 443)), (2, 1, 6, "", ("127.0.0.1", 443))],
    )
    with pytest.raises(ValueError):
        public_addresses("https://example.test")
    with pytest.raises(ValueError):
        public_addresses("https://user:password@example.test")
    with pytest.raises(ValueError):
        public_addresses("file:///etc/passwd")


def test_login_interstitial_is_not_treated_as_research_content():
    with pytest.raises(ValueError, match="未取得可用正文"):
        validate_page_body(
            "Tracxn Login", "Your browser was unable to load all resources. Login to Access Tracxn"
        )
    validate_page_body(
        "Pricing and plans", "Business is USD 20 per seat monthly. Sign in to continue managing your account."
    )
    with pytest.raises(ValueError, match="未取得可用正文"):
        validate_page_body("Help Center", "Help Center\nLoading\n×\nSorry to interrupt\nCSS Error\nRefresh")
